"""고객 목록: 엑셀(xlsx) / CSV / 구글 시트. 통화 결과를 같은 행에 다시 적는다."""
import csv
from datetime import datetime

from .config import resolve

FIELDS = ["name", "clinic", "phone", "status", "memo", "last_call"]


class CustomerBook:
    def __init__(self, cfg: dict):
        c = cfg["customers"]
        self.c = c
        self.cols = c["columns"]
        self.source = c["source"]
        if self.source == "google":
            import gspread
            gc = gspread.service_account(filename=str(resolve(c["google_credentials"])))
            sh = gc.open_by_url(c["google_sheet_url"])
            self.ws = sh.worksheet(c["sheet"]) if c.get("sheet") else sh.sheet1
            values = self.ws.get_all_values()
        elif self.source == "excel":
            import openpyxl
            self.path = resolve(c["path"])
            if not self.path.exists():  # 처음이면 열 이름만 있는 견본을 만들어 둔다
                self.path.parent.mkdir(parents=True, exist_ok=True)
                cols = c["columns"]
                head = list(dict.fromkeys([cols["name"], cols["clinic"], cols["phone"]]))
                branch = cfg.get("rules", {}).get("branch_column")
                if branch:
                    head.append(branch)
                wb = openpyxl.Workbook()
                wb.active.append(head + [cols["status"], cols["memo"], cols["last_call"]])
                wb.save(self.path)
                print(f"   고객 엑셀 견본을 만들었습니다: {self.path}  (열: {', '.join(head)})")
            self.wb = openpyxl.load_workbook(self.path)
            self.ws = self.wb[c["sheet"]] if c.get("sheet") else self.wb.worksheets[0]
            values = [["" if v is None else str(v) for v in row] for row in self.ws.iter_rows(values_only=True)]
        elif self.source == "csv":
            self.path = resolve(c["path"])
            with open(self.path, encoding="utf-8-sig", newline="") as f:
                values = list(csv.reader(f))
        else:
            raise ValueError(f"customers.source 를 excel / csv / google 중 하나로: {self.source}")
        self.header = values[0] if values else []
        self.rows = values[1:]
        for key in ("memo", "status", "last_call"):
            self._ensure_col(self.cols[key])

    def _ensure_col(self, title: str) -> None:
        if title not in self.header:
            self.header.append(title)
            self._write_cell(1, len(self.header), title)

    def _idx(self, key: str) -> int | None:
        title = self.cols.get(key, key)
        return self.header.index(title) if title in self.header else None

    def get(self, row: int) -> dict:
        """row 는 시트에서 보이는 행 번호 (헤더가 1행, 고객은 2행부터)."""
        vals = self.rows[row - 2] if 0 <= row - 2 < len(self.rows) else []
        d = {h: (vals[i] if i < len(vals) else "") for i, h in enumerate(self.header) if h}
        d["_row"] = row
        for k in FIELDS:
            i = self._idx(k)
            d["_" + k] = vals[i] if i is not None and i < len(vals) else ""
        return d

    def all(self) -> list[dict]:
        return [self.get(r) for r in range(2, len(self.rows) + 2) if any(self.rows[r - 2])]

    def next_pending(self, after: int = 1) -> dict | None:
        skip = set(self.c.get("skip_status") or [])
        for cust in self.all():
            if cust["_row"] > after and cust["_phone"] and cust["_status"] not in skip:
                return cust
        return None

    def _write_cell(self, row: int, col: int, value: str) -> None:
        if self.source == "google":
            self.ws.update_cell(row, col, value)
        elif self.source == "excel":
            self.ws.cell(row=row, column=col, value=value)
        # csv 는 save() 때 통째로 씀

    def record_result(self, row: int, status: str, memo: str) -> None:
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        old_memo = self.get(row)["_memo"]
        new_memo = f"[{stamp} AI콜] {memo}" + (f"\n{old_memo}" if old_memo else "")
        while len(self.rows) < row - 1:
            self.rows.append([])
        r = self.rows[row - 2]
        for key, val in (("status", status), ("memo", new_memo), ("last_call", stamp)):
            col = self._idx(key)
            while len(r) <= col:
                r.append("")
            r[col] = val
            self._write_cell(row, col + 1, val)
        self.save()

    def save(self) -> None:
        if self.source == "excel":
            self.wb.save(self.path)
        elif self.source == "csv":
            with open(self.path, "w", encoding="utf-8-sig", newline="") as f:
                csv.writer(f).writerows([self.header] + self.rows)
